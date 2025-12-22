from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # --- Slide 1: Title ---
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "DeepSeek OCR vs Qwen VL\n비교 분석 보고서"
    subtitle.text = "문서 파싱 성능 및 품질 평가\n2025-12-21"

    # --- Slide 2: Overview ---
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "개요 (Overview)"
    
    content = slide.placeholders[1]
    text_frame = content.text_frame
    text_frame.text = "테스트 대상 및 환경"
    
    p = text_frame.add_paragraph()
    p.text = "문서: nvme_short.pdf (총 20페이지)"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "비교 모델:"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "DeepSeek OCR (deepseek-ocr:latest)"
    p.level = 2
    
    p = text_frame.add_paragraph()
    p.text = "Qwen VL (qwen3-vl:32b-instruct-q4_K_M)"
    p.level = 2
    
    p = text_frame.add_paragraph()
    p.text = "테스트 결과: Qwen 모델의 속도 이슈로 16페이지까지 데이터 수집 후 분석"
    p.level = 1

    # --- Slide 3: Speed Comparison ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "1. 속도 비교 (Performance)"
    
    content = slide.placeholders[1]
    text_frame = content.text_frame
    text_frame.text = "DeepSeek OCR이 약 12배 더 빠름"
    
    p = text_frame.add_paragraph()
    p.text = "평균 처리 시간 (페이지당)"
    p.font.bold = True
    p.level = 0
    
    p = text_frame.add_paragraph()
    p.text = "DeepSeek: 약 11초"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "Qwen VL: 약 135초 (2분 15초)"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "전체 소요 시간 (16페이지 기준)"
    p.font.bold = True
    p.level = 0
    
    p = text_frame.add_paragraph()
    p.text = "DeepSeek: 약 3분"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "Qwen VL: 약 36분"
    p.level = 1

    # --- Slide 4: Quality Comparison ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "2. 품질 비교 (Quality)"
    
    content = slide.placeholders[1]
    text_frame = content.text_frame
    
    p = text_frame.add_paragraph()
    p.text = "DeepSeek OCR"
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 112, 192) # Blue
    
    p = text_frame.add_paragraph()
    p.text = "장점: 빠른 속도, 높은 텍스트 인식률"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "단점: 구조화 부족 (단순 텍스트 나열), 불필요한 태그 포함"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "Qwen VL"
    p.font.bold = True
    p.font.color.rgb = RGBColor(192, 0, 0) # Red
    
    p = text_frame.add_paragraph()
    p.text = "장점: 완벽한 Markdown 구조화 (목차, 표, 헤더 인식 우수)"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "단점: 매우 느린 속도, 간헐적 응답 지연(Hang)"
    p.level = 1

    # --- Slide 5: Conclusion ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "결론 및 제언 (Conclusion)"
    
    content = slide.placeholders[1]
    text_frame = content.text_frame
    text_frame.text = "사용 목적에 따른 모델 선택 권장"
    
    p = text_frame.add_paragraph()
    p.text = "대량 문서의 빠른 처리 → DeepSeek OCR"
    p.level = 1
    p.font.bold = True
    
    p = text_frame.add_paragraph()
    p.text = "텍스트 검색, 단순 인덱싱 목적에 적합"
    p.level = 2
    
    p = text_frame.add_paragraph()
    p.text = "고품질 구조화 문서 필요 → Qwen VL"
    p.level = 1
    p.font.bold = True
    
    p = text_frame.add_paragraph()
    p.text = "RAG 구축, 표/서식 보존이 필수적인 경우 적합"
    p.level = 2
    
    p = text_frame.add_paragraph()
    p.text = "단, 배치 작업으로 처리 시간을 확보해야 함"
    p.level = 2

    # Save
    output_file = "comparison_report.pptx"
    prs.save(output_file)
    print(f"Presentation saved to {output_file}")

if __name__ == "__main__":
    create_presentation()
